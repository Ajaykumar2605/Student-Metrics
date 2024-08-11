from tkinter import*
from tkinter.ttk import Combobox
from tkinter import PhotoImage
import tkinter as tk
from tkinter import messagebox
import openpyxl ,xlrd
from openpyxl import workbook
import pathlib



root=Tk()
root.title("Student Metrics")
root.geometry('700x500+300+200')
root.resizable(False,False)
root.configure(bg="#117864")

def clear():
    namevalue.set('')
    contactvalue.set('')
    agevalue.set('')
    addressentry.delete(1.0,END)

file=pathlib.Path("C:\\Users\AJAY KUMAR M\Desktop\StudentMetrics\source Files\Student_data.xlsx")
if file.exists():
    pass
else:
    file=workbook()
    sheet=file.active
    sheet['A1']="Full Name"
    sheet['B1']="Mobile No"
    sheet['D1']="Age"
    sheet['E1']="Gender"
    sheet['A1']="Address"
    file.save("Student_data.xlsx")
    
def submit():
    name=namevalue.get()
    contact=contactvalue.get()
    age=agevalue.get()
    gender=gender_combobox.get()
    address=addressentry.get(1.0,END)

    file=openpyxl.load_workbook('C:\\Users\AJAY KUMAR M\Desktop\StudentMetrics\source Files\Student_data.xlsx')
    sheet=file.active
    sheet.cell(column=1,row=sheet.max_row+1,value=name)
    sheet.cell(column=2,row=sheet.max_row,value=contact)
    sheet.cell(column=3,row=sheet.max_row,value=age)
    sheet.cell(column=4,row=sheet.max_row,value=gender)
    sheet.cell(column=5,row=sheet.max_row,value=address)

    file.save(r'C:\Users\AJAY KUMAR M\Desktop\StudentMetrics\source Files\Student_data.xlsx')

    messagebox.showinfo('info',"Detail Added! Successfully")

    namevalue.set('')
    contactvalue.set('')
    agevalue.set('')
    addressentry.delete(1.0,END)

    
#icon
icon_image = PhotoImage(file="C:\\Users\AJAY KUMAR M\Desktop\StudentMetrics\source Files\icon\Applogo.png")
root.iconphoto(False,icon_image)

#heading
Label(root,text="Please fill out this Entry form",font='arial',bg="#117864",fg='#fff').place(x=20,y=20)

#label
Label(root,text='Name:',font=23,bg="#117864",fg="#fff").place(x=50,y=100)
Label(root,text='Mobile No:',font=23,bg="#117864",fg="#fff").place(x=50,y=150)
Label(root,text='Age:',font=23,bg="#117864",fg="#fff").place(x=50,y=200)
Label(root,text='Gender:',font=23,bg="#117864",fg="#fff").place(x=370,y=200)
Label(root,text='Address',font=23,bg="#117864",fg="#fff").place(x=50,y=250)

#Entry
namevalue=StringVar()
contactvalue=StringVar()
agevalue=StringVar()

nameentry=Entry(root,textvariable=namevalue,width=45,bd=2,font=20)
contactentry=Entry(root,textvariable=contactvalue,width=45,bd=2,font=20)
ageentry=Entry(root,textvariable=agevalue,width=15,bd=2,font=20)
#gender
gender_combobox = Combobox(root,values=['Male','Female'],font="arial, 14",state='r',width=14)
gender_combobox.place(x=440,y=200)
gender_combobox.set('Male')
#address
addressentry=Text(root,width=50,height=4,bd=4)

nameentry.place(x=200,y=100)
contactentry.place(x=200,y=150)
ageentry.place(x=200,y=200)
addressentry.place(x=200,y=250)
#button
Button(root,text="Submit",bg="#717d7e",fg="White",width=15,height=2,command=submit).place(x=200,y=350)
Button(root,text="Clear",bg="#717d7e",fg="White",width=15,height=2,command=clear).place(x=340,y=350)
Button(root,text="Exit",bg="#717d7e",fg="White",width=15,height=2,command=lambda:root.destroy()).place(x=480,y=350)
#copy right logo
footer_frame = tk.Frame(root, pady=10)
footer_frame.pack(side=tk.BOTTOM, fill=tk.X)
copyright_notice = tk.Label(footer_frame, text="© 2024 Ajay Kumar. All rights reserved.", anchor='e')
copyright_notice.pack(side=tk.BOTTOM, padx=10)

root.mainloop()
